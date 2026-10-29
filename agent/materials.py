"""Convierte lo que adjunta la directora (documentos, fotos, enlaces) en material para el guionista."""
from __future__ import annotations

import base64
import re
import subprocess
from pathlib import Path

from .common import log
from .news import fetch_article_text

MAX_TEXT = 15000
MAX_IMAGES = 5
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".heic"}
URL_RE = re.compile(r"https?://[^\s<>\"')\]]+")


def _pdf(p: Path) -> str:
    from pypdf import PdfReader
    return "\n".join((page.extract_text() or "") for page in PdfReader(str(p)).pages[:40])


def _docx(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    rows = [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
    return "\n".join([para.text for para in d.paragraphs] + rows)


def _pptx(p: Path) -> str:
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(str(p)).slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        out.append(f"[Diapositiva {i}] " + " / ".join(texts))
    return "\n".join(out)


def _xlsx(p: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(p), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets[:5]:
        out.append(f"[Hoja {ws.title}]")
        for row in ws.iter_rows(max_row=150, values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _image(p: Path) -> dict | None:
    """Imagen lista para Claude (se reduce a 1568 px y JPEG para no pasarse de tamaño)."""
    small = p.with_name(p.stem + "_claude.jpg")
    try:
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", str(p), "-vf",
                        "scale='min(1568,iw)':'min(1568,ih)':force_original_aspect_ratio=decrease",
                        "-frames:v", "1", "-q:v", "4", str(small)], check=True)
    except subprocess.CalledProcessError:
        log.warning("No se pudo leer la imagen %s", p.name)
        return None
    return {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                         "data": base64.b64encode(small.read_bytes()).decode()}}


READERS = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx}


def extract(files: list[Path], urls: list[str]) -> tuple[str, list[dict]]:
    """Devuelve (texto con todo el material, imágenes para Claude)."""
    parts, images = [], []
    for f in files:
        ext = f.suffix.lower()
        try:
            if ext in IMAGE_EXT:
                if len(images) < MAX_IMAGES and (img := _image(f)):
                    images.append(img)
                    parts.append(f"### Imagen adjunta: {f.name} (se te muestra aparte)")
                continue
            if ext in READERS:
                text = READERS[ext](f)
            elif ext in {".txt", ".md", ".csv"}:
                text = f.read_text(encoding="utf-8", errors="replace")
            else:
                log.info("Tipo de archivo no soportado, se ignora: %s", f.name)
                continue
        except Exception as e:  # un archivo raro no debe tumbar el vídeo
            log.warning("No se pudo leer %s: %s", f.name, e)
            continue
        text = re.sub(r"[ \t]+", " ", text).strip()
        if text:
            parts.append(f"### Documento adjunto: {f.name}\n{text[:8000]}")
    for u in urls[:5]:
        text = fetch_article_text(u)
        if text:
            parts.append(f"### Enlace: {u}\n{text}")
    return "\n\n".join(parts)[:MAX_TEXT], images


def urls_in(text: str) -> list[str]:
    return [u.rstrip(".,;") for u in URL_RE.findall(text or "") if "mailto:" not in u]
